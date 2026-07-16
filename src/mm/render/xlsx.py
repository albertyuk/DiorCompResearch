"""Companion xlsx dump of the project table."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from ..dates import date_text, parse_iso
from .deck import BrandSpec, _social_text


XLSX_FONT = "Futura Lt BT"    # owner spec: spreadsheet is Futura Lt BT 11
XLSX_SIZE = 11


def write_projects_xlsx(brands: list[BrandSpec], out_path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "PROJECTS"
    headers = ["BRAND", "DATE", "PROJECT", "ASSETS", "SOCIAL", "EC", "AD",
               "PHASE", "ONGOING"]
    ws.append(headers)
    head_fill = PatternFill("solid", fgColor="000000")
    for c in ws[1]:
        c.font = Font(name=XLSX_FONT, size=XLSX_SIZE, bold=True, color="FFFFFF")
        c.fill = head_fill
        c.alignment = Alignment(horizontal="center")
    for b in brands:
        for p in b.projects:
            ws.append([
                b.display_name,
                date_text(parse_iso(p.date_start), parse_iso(p.date_end), p.ongoing),
                p.description, p.assets, _social_text(p.platforms), "N/A", "N/A",
                p.phase_suffix or "", "YES" if p.ongoing else "",
            ])
            for c in ws[ws.max_row]:
                c.font = Font(name=XLSX_FONT, size=XLSX_SIZE)
    widths = [16, 18, 70, 14, 28, 6, 6, 14, 9]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out_path))
    return out_path
