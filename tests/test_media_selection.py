"""HQ images: max-resolution variant picking, review-time image selection,
manual HQ uploads, selected-images-win rendering, and the owner font specs
(xlsx Futura Lt BT 11 · celeb labels Calibre 9)."""
import json

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import select

PASS = "team-pass-123"
SECRET = "f" * 64
PNG = b"\x89PNG\r\n\x1a\n" + b"x" * 32


@pytest.fixture
def tmp_db(tmp_path, monkeypatch):
    import mm.db as mdb
    monkeypatch.setattr(mdb, "DB_PATH", tmp_path / "t.db")
    monkeypatch.setattr(mdb, "_engine", None)
    import mm.media as mmedia
    monkeypatch.setattr(mmedia, "RUNS_DIR", tmp_path / "runs")
    yield mdb
    if mdb._engine is not None:
        mdb._engine.dispose()


@pytest.fixture
def client(tmp_db, monkeypatch):
    monkeypatch.setenv("CONSOLE_PASSPHRASE", PASS)
    monkeypatch.setenv("MM_SECRET_KEY", SECRET)
    from mm.console import create_app
    c = TestClient(create_app())
    r = c.post("/login", data={"name": "Alice", "passphrase": PASS},
               follow_redirects=False)
    assert r.status_code == 303
    return c


def _seed(mdb, post_id="weibo:M1", media=None):
    with mdb.get_engine().begin() as conn:
        mdb.upsert(conn, mdb.posts, {
            "post_id": post_id, "month": "2026-06", "brand": "lv",
            "platform": "weibo", "url": "https://weibo.com/1/x",
            "created_at": "2026-06-05T12:00:00+08:00", "caption": "上海快闪 王一博",
            "at_tags": "[]", "hashtags": "[]",
            "media": json.dumps(media or [], ensure_ascii=False),
            "is_repost": False, "repost_ambiguous": False}, ["post_id"])


# ── max-resolution variant ───────────────────────────────────────────────────

def test_largest_pic_variant_picks_max_area():
    from mm.normalize import _largest_pic_variant
    info = {"thumbnail": {"url": "t", "width": 120, "height": 120},
            "large": {"url": "l", "width": 1080, "height": 1350},
            "original": {"url": "o", "width": 2048, "height": 2560},
            "largest": {"url": "L", "width": 4096, "height": 5120}}
    assert _largest_pic_variant(info)["url"] == "L"
    # dimensions beat the preference chain when a bigger variant exists
    info["mw4000"] = {"url": "XL", "width": 6000, "height": 8000}
    assert _largest_pic_variant(info)["url"] == "XL"
    # unknown names without dimensions still yield a url
    assert _largest_pic_variant({"weird": {"url": "w"}})["url"] == "w"
    assert _largest_pic_variant({}) == {}


# ── selection + upload routes ────────────────────────────────────────────────

def test_media_select_toggle_and_validation(client, tmp_db, tmp_path):
    img = tmp_path / "a.jpg"
    img.write_bytes(b"\xff\xd8\xff" + b"j" * 16)
    _seed(tmp_db, media=[{"kind": "image", "url": "u", "local_path": str(img)}])
    r = client.post("/review/2026-06/posts/weibo:M1/media/select",
                    data={"idx": 0, "selected": "true"})
    assert r.status_code == 200 and r.json()["selected"] is True
    with tmp_db.get_engine().connect() as conn:
        media = json.loads(conn.execute(select(tmp_db.posts.c.media)).scalar())
        row = tmp_db.last_audit(conn, "media_select", "post", "weibo:M1#0")
    assert media[0]["selected"] is True and row["actor_name"] == "Alice"
    r = client.post("/review/2026-06/posts/weibo:M1/media/select",
                    data={"idx": 0, "selected": "false"})
    assert r.json()["selected"] is False
    assert client.post("/review/2026-06/posts/weibo:M1/media/select",
                       data={"idx": 9, "selected": "true"}).status_code == 400
    assert client.post("/review/2026-06/posts/weibo:GHOST/media/select",
                       data={"idx": 0, "selected": "true"}).status_code == 404


def test_hq_upload_stores_selected_media(client, tmp_db):
    _seed(tmp_db)
    r = client.post("/review/2026-06/posts/weibo:M1/upload",
                    files={"files": ("orig.png", PNG, "image/png")})
    assert r.status_code == 200 and r.json()["added"] == 1
    with tmp_db.get_engine().connect() as conn:
        media = json.loads(conn.execute(select(tmp_db.posts.c.media)).scalar())
    assert len(media) == 1
    m = media[0]
    assert m["source"] == "upload" and m["selected"] is True
    assert m["local_path"].endswith(".png")
    from pathlib import Path
    assert Path(m["local_path"]).read_bytes() == PNG
    # duplicate content is not appended twice
    client.post("/review/2026-06/posts/weibo:M1/upload",
                files={"files": ("again.png", PNG, "image/png")})
    with tmp_db.get_engine().connect() as conn:
        media = json.loads(conn.execute(select(tmp_db.posts.c.media)).scalar())
    assert len(media) == 1
    # non-image rejected
    r = client.post("/review/2026-06/posts/weibo:M1/upload",
                    files={"files": ("evil.png", b"#!/bin/sh", "image/png")})
    assert r.status_code == 400


# ── renderer: selected images win ────────────────────────────────────────────

def test_project_visuals_prefer_selected_images(tmp_db, tmp_path):
    from mm.pipeline import _project_visuals
    sel1, sel2, unsel = (tmp_path / n for n in ("s1.jpg", "s2.jpg", "u.jpg"))
    for f in (sel1, sel2, unsel):
        f.write_bytes(b"\xff\xd8\xff")
    _seed(tmp_db, post_id="weibo:V1", media=[
        {"kind": "image", "local_path": str(sel1), "selected": True},
        {"kind": "image", "local_path": str(unsel)},
        {"kind": "image", "local_path": str(sel2), "selected": True}])
    _seed(tmp_db, post_id="weibo:V2", media=[])          # no selection
    with tmp_db.get_engine().begin() as conn:
        pid = conn.execute(tmp_db.projects.insert().values(
            month="2026-06", brand="lv", title="T", status="confirmed",
            assets="PHOTO", celebs="[]", hero_media="[]")).inserted_primary_key[0]
        for post_id in ("weibo:V1", "weibo:V2"):
            conn.execute(tmp_db.project_posts.insert().values(
                project_id=pid, post_id=post_id))

    class FakeFactory:
        def visual_for_post(self, brand, post):
            return tmp_path / "card.png"

    celebs = [{"name_cn": "王一博", "relation_display": "BRAND AMBASSADOR",
               "display": "WANG YIBO"}]
    with tmp_db.get_engine().connect() as conn:
        vis = _project_visuals(conn, FakeFactory(), "lv",
                               {"id": pid, "assets": "PHOTO"}, celebs)
    images = [v["image"] for v in vis]
    assert str(sel1) in images and str(sel2) in images   # both ticked images
    assert str(unsel) not in images                      # unticked skipped
    assert str(tmp_path / "card.png") in images          # V2 fell back to card
    # celeb label rides on the FIRST selected image of the post only
    v1 = [v for v in vis if v["image"] == str(sel1)][0]
    v2 = [v for v in vis if v["image"] == str(sel2)][0]
    assert v1["label_name"] == "WANG YIBO" and v2["label_name"] is None


# ── owner font specs ─────────────────────────────────────────────────────────

def test_xlsx_font_is_futura_lt_bt_11(tmp_path):
    from openpyxl import load_workbook
    from mm.render.deck import BrandSpec, ProjectSpec
    from mm.render.xlsx import write_projects_xlsx
    spec = [BrandSpec(key="lv", display_name="LOUIS VUITTON", projects=[
        ProjectSpec(title="T", phase_suffix=None, date_start="2026-06-05",
                    date_end=None, ongoing=False, assets="PHOTO",
                    platforms=["weibo"], description="T", visuals=[])])]
    out = write_projects_xlsx(spec, tmp_path / "t.xlsx")
    ws = load_workbook(out).active
    for cell in ("A1", "C1", "A2", "C2"):
        assert ws[cell].font.name == "Futura Lt BT", cell
        assert ws[cell].font.size == 11, cell


def test_celeb_labels_are_calibre_9(tmp_path):
    from pptx import Presentation
    from mm.config import ROOT
    from mm.render.deck import BrandSpec, DeckBuilder, ProjectSpec, Visual
    icon = ROOT / "template" / "icons" / "weibo.png"
    spec = [BrandSpec(key="lv", display_name="LOUIS VUITTON", projects=[
        ProjectSpec(title="T", phase_suffix=None, date_start="2026-06-05",
                    date_end=None, ongoing=False, assets="PHOTO",
                    platforms=["weibo"], description="T",
                    visuals=[Visual(image=str(icon), kind="photo", link=None,
                                    label_top="BRAND AMBASSADOR",
                                    label_name="WANG YIBO")])])]
    out = tmp_path / "t.pptx"
    DeckBuilder().build(spec, out)
    found = {}
    for slide in Presentation(str(out)).slides:
        for sh in slide.shapes:
            if not sh.has_text_frame:
                continue
            for para in sh.text_frame.paragraphs:
                for r in para.runs:
                    if r.text in ("BRAND AMBASSADOR", "WANG YIBO"):
                        found[r.text] = (r.font.size.pt, r.font.name)
    assert found["BRAND AMBASSADOR"] == (9, "Calibre")
    assert found["WANG YIBO"] == (9, "Calibre")
