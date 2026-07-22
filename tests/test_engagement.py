"""Engagement figures (owner request): likes/comments/shares/views per post —
extracted at ingest from every platform's payload, backfilled for free from
the archived raw pages, shown on both review pages, summed per project in
the XLSX. Weibo exposes no view count (verified on live payloads); the
numbers are a snapshot from ingest time.
"""
import json

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import select

PASS = "team-pass-123"


@pytest.fixture
def tmp_db(tmp_path, monkeypatch):
    import mm.db as mdb
    monkeypatch.setattr(mdb, "DB_PATH", tmp_path / "t.db")
    monkeypatch.setattr(mdb, "_engine", None)
    import mm.media as mmedia
    monkeypatch.setattr(mmedia, "RUNS_DIR", tmp_path / "runs")
    yield mdb
    if mdb._engine is not None:
        mdb._engine.dispose()


@pytest.fixture
def client(tmp_db, monkeypatch):
    monkeypatch.setenv("CONSOLE_PASSPHRASE", PASS)
    monkeypatch.setenv("MM_SECRET_KEY", "f" * 64)
    from mm.console import create_app
    c = TestClient(create_app())
    r = c.post("/login", data={"name": "Alice", "passphrase": PASS},
               follow_redirects=False)
    assert r.status_code == 303
    return c


def _post(mdb, post_id, *, engagement=None, brand="lv", platform="weibo",
          keep=True):
    with mdb.get_engine().begin() as conn:
        mdb.upsert(conn, mdb.posts, {
            "post_id": post_id, "month": "2026-06", "brand": brand,
            "platform": platform, "url": f"https://x/{post_id}",
            "created_at": "2026-06-05T12:00:00+08:00", "caption": "上海快闪",
            "at_tags": "[]", "hashtags": "[]", "media": "[]",
            "is_repost": False, "repost_ambiguous": False,
            "engagement": json.dumps(engagement) if engagement is not None
            else None}, ["post_id"])
        mdb.upsert(conn, mdb.verdicts, {
            "post_id": post_id, "keep": keep, "confidence": 0.9,
            "reasons": "[]", "celebs_tagged": "[]", "category": "event",
            "media_focus": "photo", "needs_review": False}, ["post_id"])


def test_count_parses_ints_and_cn_display_strings():
    from mm.normalize import _count
    assert _count(19) == 19
    assert _count(0) == 0
    assert _count("1,234") == 1234
    assert _count("1.2万") == 12000
    assert _count("10万+") == 100000
    assert _count("3亿") == 300000000
    assert _count("4.5w") == 45000
    assert _count(None) is None
    assert _count("赞") is None
    assert _count(True) is None
    assert _count(-3) is None


def test_normalizers_extract_engagement_per_platform():
    from mm import normalize
    wb = normalize.normalize_weibo(
        {"mblogid": "M1", "created_at": "Fri Jul 04 12:00:00 +0800 2026",
         "text_raw": "hi", "attitudes_count": 19, "comments_count": 1,
         "reposts_count": 4, "favorites_count": 0}, "42")
    assert wb["engagement"] == {"likes": 19, "comments": 1, "shares": 4,
                                "favorites": 0}
    xhs = normalize.normalize_xhs(
        {"note_id": "N1", "desc": "d", "interact_info": {
            "liked_count": "1.2万", "comment_count": "88",
            "collected_count": "456"}})
    assert xhs["engagement"] == {"likes": 12000, "comments": 88,
                                 "favorites": 456}
    dy = normalize.normalize_douyin(
        {"aweme_id": "A1", "desc": "d", "statistics": {
            "digg_count": 100, "comment_count": 5, "share_count": 7,
            "play_count": 90000, "collect_count": 3}})
    assert dy["engagement"] == {"likes": 100, "comments": 5, "shares": 7,
                                "views": 90000, "favorites": 3}
    mp = normalize.normalize_wechat_mp(
        {"appmsgid": "W1", "title": "t", "link": "https://x",
         "read_num": 100001, "like_num": 88})
    assert mp["engagement"] == {"views": 100001, "likes": 88}
    # payloads without stats → empty dict, never a crash
    assert normalize.normalize_weibo(
        {"mblogid": "M2", "text_raw": "x"}, "42")["engagement"] == {}


def test_migration_adds_engagement_to_an_old_db(tmp_path, monkeypatch):
    import sqlite3

    import mm.db as mdb
    db_file = tmp_path / "old.db"
    con = sqlite3.connect(db_file)
    con.execute("CREATE TABLE posts (post_id VARCHAR PRIMARY KEY, "
                "month VARCHAR, brand VARCHAR, platform VARCHAR)")
    con.commit()
    con.close()
    monkeypatch.setattr(mdb, "DB_PATH", db_file)
    monkeypatch.setattr(mdb, "_engine", None)
    eng = mdb.get_engine()
    with eng.connect() as c:
        cols = [r[1] for r in c.exec_driver_sql("PRAGMA table_info(posts)")]
    assert "engagement" in cols
    eng.dispose()


def test_backfill_engagement_reparses_raw_archives(tmp_db, tmp_path):
    """Months ingested before the column existed get numbers for FREE: the
    per-post raw_path points at the archived API page on disk."""
    from mm.ingest import backfill_engagement
    raw = tmp_path / "weibo_page001.json"
    raw.write_text(json.dumps({"data": {"items": [
        {"category": "post", "data": {
            "mblogid": "OLD1", "text_raw": "hello",
            "created_at": "Fri Jul 04 12:00:00 +0800 2026",
            "attitudes_count": 19, "comments_count": 1,
            "reposts_count": 4}}]}}, ensure_ascii=False))
    _post(tmp_db, "weibo:OLD1")
    with tmp_db.get_engine().begin() as conn:
        conn.execute(tmp_db.posts.update()
                     .where(tmp_db.posts.c.post_id == "weibo:OLD1")
                     .values(raw_path=str(raw), engagement=None))
    assert backfill_engagement(tmp_db.get_engine()) == 1
    with tmp_db.get_engine().connect() as conn:
        eng = json.loads(conn.execute(
            select(tmp_db.posts.c.engagement)
            .where(tmp_db.posts.c.post_id == "weibo:OLD1")).scalar())
    assert eng == {"likes": 19, "comments": 1, "shares": 4}
    # idempotent: nothing left to fill on the second run
    assert backfill_engagement(tmp_db.get_engine()) == 0


def test_fmt_count_and_eng_line_localize():
    from mm.console import _eng_line, _fmt_count
    assert _fmt_count(999, "en") == "999"
    assert _fmt_count(19_000, "en") == "19k"
    assert _fmt_count(1_230_000, "en") == "1.2M"
    assert _fmt_count(19_000, "zh") == "1.9万"
    assert _fmt_count(300_000_000, "zh") == "3亿"
    raw = json.dumps({"likes": 19000, "comments": 123, "shares": 45})
    assert _eng_line(raw, "en") == "likes 19k · cmts 123 · shares 45"
    assert _eng_line(raw, "zh") == "赞 1.9万 · 评 123 · 转 45"
    assert _eng_line("{}", "en") is None
    assert _eng_line(None, "en") is None
    assert _eng_line("not json", "en") is None


def test_review_pages_show_the_engagement_line(client, tmp_db):
    _post(tmp_db, "weibo:E1", engagement={"likes": 19000, "comments": 123})
    page = client.get("/review/2026-06/posts").text
    assert "likes 19k" in page and "cmts 123" in page
    # review 2: the member row carries the same line
    with tmp_db.get_engine().begin() as conn:
        pid = conn.execute(tmp_db.projects.insert().values(
            month="2026-06", brand="lv", title="ENG", phase_suffix=None,
            date_start="2026-06-05", date_end="2026-06-05", ongoing=False,
            assets="PHOTO", description="ENG", celebs="[]", hero_media="[]",
            status="draft")).inserted_primary_key[0]
        conn.execute(tmp_db.project_posts.insert().values(
            project_id=pid, post_id="weibo:E1", role="member"))
    page2 = client.get("/review/2026-06/projects").text
    assert "likes 19k" in page2


def test_xlsx_sums_member_engagement_per_project(tmp_path):
    from openpyxl import load_workbook

    from mm.render.deck import BrandSpec, ProjectSpec
    from mm.render.xlsx import write_projects_xlsx
    spec = [BrandSpec(key="lv", display_name="LOUIS VUITTON", projects=[
        ProjectSpec(title="SHOW", phase_suffix=None, date_start="2026-06-05",
                    date_end=None, ongoing=False, assets="PHOTO",
                    platforms=["weibo"], description="SHOW",
                    engagement={"likes": 31000, "comments": 124,
                                "shares": 49}),
        ProjectSpec(title="NO DATA", phase_suffix=None,
                    date_start="2026-06-06", date_end=None, ongoing=False,
                    assets="PHOTO", platforms=["weibo"],
                    description="NO DATA")])]
    out = tmp_path / "p.xlsx"
    write_projects_xlsx(spec, out)
    ws = load_workbook(out).active
    head = [c.value for c in ws[1]]
    assert head[9:13] == ["LIKES", "COMMENTS", "SHARES", "VIEWS"]
    assert [c.value for c in ws[2]][9:12] == [31000, 124, 49]
    assert [c.value for c in ws[3]][9:13] == [None, None, None, None]
