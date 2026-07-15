import json
from pathlib import Path

from mm import normalize
from mm.render.deck import (GRID_CAP, DeckBuilder, _social_text, load_spec)

ROOT = Path(__file__).resolve().parents[1]


# -- weibo normalization ------------------------------------------------------------

def _mblog(**over):
    base = {
        "mblogid": "Ptest123", "created_at": "Wed Jul 08 12:30:00 +0800 2026",
        "text_raw": "品牌大使@张三 亮相活动 #活动#",
        "pic_ids": ["p1"],
        "pic_infos": {"p1": {"original": {"url": "https://wx1.sinaimg.cn/large/p1.jpg"}}},
        "user": {"screen_name": "香奈儿CHANEL", "avatar_hd": "https://a/x.jpg"},
    }
    base.update(over)
    return base


def test_normalize_weibo_basic():
    p = normalize.normalize_weibo(_mblog(), "1892475055")
    assert p["post_id"] == "weibo:Ptest123"
    assert p["url"] == "https://weibo.com/1892475055/Ptest123"
    assert p["at_tags"] == ["张三"]
    assert p["hashtags"] == ["活动"]
    assert p["media"][0]["url"].endswith("p1.jpg")
    assert p["created_at"].startswith("2026-07-08")
    assert not p["is_repost"]


def test_normalize_weibo_pure_repost_flagged():
    p = normalize.normalize_weibo(
        _mblog(text_raw="转发微博", retweeted_status={"id": 1}), "1")
    assert p["is_repost"] and not p["repost_ambiguous"]


def test_normalize_weibo_commented_repost_ambiguous():
    p = normalize.normalize_weibo(
        _mblog(text_raw="这个系列真的很好看，值得一看//@某人:原文",
               retweeted_status={"id": 1}), "1")
    assert not p["is_repost"] and p["repost_ambiguous"]


def test_normalize_weibo_video_cover():
    p = normalize.normalize_weibo(_mblog(
        pic_ids=[], pic_infos={},
        page_info={"object_type": "video",
                   "page_pic": {"url": "https://wx1.sinaimg.cn/cover.jpg"}}), "1")
    assert p["media"][0]["kind"] == "video_cover"


def test_find_post_list_nested():
    payload = {"code": 200, "data": {"data": {"list": [
        {"mblogid": "a", "text_raw": "x"}, {"mblogid": "b", "text_raw": "y"}]}}}
    assert len(normalize.weibo_posts_from_response(payload)) == 2


# -- deck building -------------------------------------------------------------------

def test_social_text_order_and_wechat_merge():
    assert _social_text(["wechat_channels", "xhs", "weibo"]) == "WEIBO/RED/WECHAT"


def test_fixture_deck_builds(tmp_path):
    spec = json.loads((ROOT / "fixtures/projects.json").read_text())
    brands = load_spec(spec)
    out = DeckBuilder(ROOT / "template/reference.pptx").build(
        brands, tmp_path / "out.pptx")
    from pptx import Presentation
    prs = Presentation(str(out))
    # cover + stub + per brand (table + project slides)
    expected = 2 + sum(1 + len(b.projects) for b in brands)
    assert len(prs.slides) == expected
    # first two slides survived the clone
    assert "COMPETITOR LOCAL ASSETS" in prs.slides[0].shapes[0].text_frame.text


def test_grid_overflow_creates_continuation(tmp_path):
    spec = json.loads((ROOT / "fixtures/projects.json").read_text())
    brands = load_spec(spec)
    fendi = next(b for b in brands if b.key == "fendi")
    baguette = fendi.projects[0]
    baguette.visuals = baguette.visuals * 2   # 28 visuals → 2 slides
    b = DeckBuilder(ROOT / "template/reference.pptx")
    slides = b.add_project_slides(fendi, baguette)
    assert len(slides) == 2
    assert len(baguette.visuals) == 2 * GRID_CAP


def test_short_commentary_repost_is_ambiguous_not_dropped():
    p = normalize.normalize_weibo(
        _mblog(text_raw="好看！", retweeted_status={"id": 1}), "1")
    assert not p["is_repost"] and p["repost_ambiguous"]


def test_video_project_chunks_by_six():
    spec = json.loads((ROOT / "fixtures/projects.json").read_text())
    brands = load_spec(spec)
    chanel = next(b for b in brands if b.key == "chanel")
    video = chanel.projects[0]
    assert video.assets == "VIDEO"
    video.visuals = video.visuals * 7   # 14 stills → 3 slides of ≤6
    b = DeckBuilder(ROOT / "template/reference.pptx")
    slides = b.add_project_slides(chanel, video)
    assert len(slides) == 3


def test_missing_visual_file_skipped_not_fatal(tmp_path):
    spec = json.loads((ROOT / "fixtures/projects.json").read_text())
    brands = load_spec(spec)
    gucci = next(b for b in brands if b.key == "gucci")
    gucci.projects[0].visuals[0].image = "fixtures/media/DOES_NOT_EXIST.jpg"
    out = DeckBuilder(ROOT / "template/reference.pptx").build(
        [gucci], tmp_path / "g.pptx")
    assert out.exists()


def test_table_splits_after_seven_projects(tmp_path):
    spec = json.loads((ROOT / "fixtures/projects.json").read_text())
    brands = load_spec(spec)
    fendi = next(b for b in brands if b.key == "fendi")
    fendi.projects = fendi.projects * 5      # 10 projects → 2 table slides
    for p in fendi.projects:
        p.visuals = []
    b = DeckBuilder(ROOT / "template/reference.pptx")
    slides = b.add_table_slides(fendi)
    assert len(slides) == 2
